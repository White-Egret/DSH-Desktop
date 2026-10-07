# DSH Desktop —— Python 能力桥接（生成文件，请勿手工编辑）

"""Word 工具：读 .docx 内容 + 生成带样式的报告。

读取走 python-docx 的段落/表格 API（结构化，不含样式）；
需要「尽可能忠实」的原文抽取时，模型应该改用 convert_file_to_markdown。
"""

from __future__ import annotations

import dsh_bridge

from utils import DEFAULT_MAX_CHARS, ToolError, resolve_input_path, run_tool


@dsh_bridge.tool(
    name="read_docx_text",
    description=(
        "读取 Word 文档（.docx）的正文段落与表格，返回结构化 JSON。"
        "当用户给了 .docx 让你「读内容 / 看表格 / 提取章节」时用它。"
        "返回 {paragraphs, tables, chars}；字符数超过 max_chars 时正文会截断并标明。"
        "若需要保留原始排版转成 Markdown，请改用 convert_file_to_markdown。"
    ),
    parameters={
        "type": "object",
        "properties": {
            "file_path": {"type": "string", "description": "本机 .docx 文件绝对路径"},
            "include_tables": {
                "type": "boolean",
                "description": "是否一并返回表格内容，默认 true",
            },
            "max_chars": {
                "type": "integer",
                "description": f"正文总长度上限，超过截断；默认 {DEFAULT_MAX_CHARS}",
                "minimum": 0,
            },
        },
        "required": ["file_path"],
    },
)
def read_docx_text(
    file_path: str,
    include_tables: bool = True,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> dict:
    def work():
        try:
            from docx import Document
        except ImportError as exc:  # noqa: BLE001
            raise ToolError(f"python-docx 不可用：{exc}。请重新执行「基本安装」。") from exc

        src = resolve_input_path(file_path)
        try:
            doc = Document(str(src))
        except Exception as exc:  # noqa: BLE001 — .doc 老格式 / 加密 / 损坏
            raise ToolError(
                f"打开 {src.name} 失败：{type(exc).__name__}: {exc}"
                "（注意：python-docx 只支持 .docx，不支持老的 .doc）"
            ) from exc

        limit = max_chars if isinstance(max_chars, int) and max_chars > 0 else DEFAULT_MAX_CHARS
        paragraphs: list[str] = []
        used = 0
        truncated = False
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if used + len(text) > limit:
                truncated = True
                break
            paragraphs.append(text)
            used += len(text)

        out: dict = {
            "source": str(src),
            "paragraphs": paragraphs,
            "chars": used,
            "truncated": truncated,
        }
        if include_tables:
            tables = []
            for t in doc.tables:
                tables.append([[c.text.strip() for c in row.cells] for row in t.rows])
            out["tables"] = tables
            out["table_count"] = len(tables)
        return out

    return run_tool(work)


@dsh_bridge.tool(
    name="generate_word_report",
    description=(
        "生成一份带样式的 Word 报告（标题 + 小节 + 可选表格）。"
        "适合「把分析结论整理成一份可发出去的 Word」。"
        "sections 是 [{heading, paragraphs: [...]}]，title 是文档大标题。"
        "table 可选：{headers: [...], rows: [[...]]} 会追加在最后。"
        "返回 {path, paragraphs_written, tables_written}。"
    ),
    parameters={
        "type": "object",
        "properties": {
            "output_path": {"type": "string", "description": "要写出的 .docx 绝对路径"},
            "title": {"type": "string", "description": "文档大标题"},
            "sections": {
                "type": "array",
                "description": "小节列表 [{heading: 小节标题, paragraphs: [段落...]}]",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string"},
                        "paragraphs": {"type": "array", "items": {"type": "string"}},
                    },
                },
            },
            "table": {
                "type": "object",
                "description": '可选表格 {headers: [...], rows: [[...]]}',
            },
        },
        "required": ["output_path", "title"],
    },
)
def generate_word_report(
    output_path: str,
    title: str,
    sections=None,
    table: dict | None = None,
) -> dict:
    def work():
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError as exc:  # noqa: BLE001
            raise ToolError(f"python-docx 不可用：{exc}。请重新执行「基本安装」。") from exc

        if not isinstance(sections, (list, tuple)):
            sections = []
        dest = resolve_input_path(output_path, must_exist=False)
        if dest.suffix.lower() != ".docx":
            dest = dest.with_suffix(".docx")
        dest.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        # 中文字体：Word 默认主题字体对 CJK 不友好，显式指定东亚字体才不至于串行
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        try:
            style.element.rPr.rFonts.set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
                "微软雅黑",
            )
        except Exception:  # noqa: BLE001 — 老版本 docx 的 rPr 可能为 None
            pass

        doc.add_heading(str(title), level=0)
        written = 0
        for sec in sections:
            if not isinstance(sec, dict):
                continue
            heading = sec.get("heading")
            if heading:
                doc.add_heading(str(heading), level=1)
            for para in sec.get("paragraphs") or []:
                doc.add_paragraph(str(para))
                written += 1

        tables_written = 0
        if isinstance(table, dict) and table.get("rows"):
            headers = [str(h) for h in (table.get("headers") or [])]
            rows = table["rows"]
            n_cols = max([len(headers)] + [len(r) for r in rows if isinstance(r, (list, tuple))] or [0])
            n_cols = max(n_cols, 1)
            tbl = doc.add_table(rows=1 + len(rows), cols=n_cols)
            tbl.style = "Table Grid"  # 没有边框的表格在打印出来时是一堆空白
            for ci in range(n_cols):
                tbl.cell(0, ci).text = headers[ci] if ci < len(headers) else ""
            for ri, row in enumerate(rows, start=1):
                if not isinstance(row, (list, tuple)):
                    row = [row]
                for ci in range(n_cols):
                    tbl.cell(ri, ci).text = str(row[ci]) if ci < len(row) else ""
            tables_written = 1
            written += 1

        doc.save(str(dest))
        return {
            "path": str(dest),
            "paragraphs_written": written,
            "tables_written": tables_written,
        }

    return run_tool(work)