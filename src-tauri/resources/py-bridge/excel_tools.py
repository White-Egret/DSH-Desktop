# DSH Desktop —— Python 能力桥接（生成文件，请勿手工编辑）

"""Excel 工具：openpyxl 读写 + XlsxWriter 写带样式的表。

刻意**不重复造轮子**：DSH 自带的文件工具能读写文本行，但 xlsx 是 zip + XML 的
二进制格式，文本工具碰不了 —— 所以这里是真正的差异化能力。
"""

from __future__ import annotations

import dsh_bridge

from utils import (
    DEFAULT_MAX_CHARS,
    ToolError,
    resolve_input_path,
    run_tool,
)


@dsh_bridge.tool(
    name="read_excel_data",
    description=(
        "读取 Excel（.xlsx / .xlsm）里一个工作表的内容，返回行与列的 JSON。"
        "当用户给了表格让你「看数据 / 统计 / 取某几列」时用它。"
        "返回 {sheet, sheets, headers, rows, row_count, truncated}；"
        "行数超过 max_rows 时只返回前 max_rows 行并标明 truncated。"
    ),
    parameters={
        "type": "object",
        "properties": {
            "file_path": {"type": "string", "description": "本机 .xlsx/.xlsm 文件绝对路径"},
            "sheet_name": {
                "type": "string",
                "description": "工作表名；留空 = 第一个工作表",
            },
            "max_rows": {
                "type": "integer",
                "description": "最多返回多少行，默认 200",
                "minimum": 1,
            },
        },
        "required": ["file_path"],
    },
)
def read_excel_data(file_path: str, sheet_name: str | None = None, max_rows: int = 200) -> dict:
    def work():
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # noqa: BLE001
            raise ToolError(f"openpyxl 不可用：{exc}。请重新执行「基本安装」。") from exc

        src = resolve_input_path(file_path)
        limit = max_rows if isinstance(max_rows, int) and max_rows > 0 else 200
        try:
            # data_only=True：读到的是**公式算出来的值**，而不是公式本身
            # （用户问「这个数是多少」时，公式字符串不是答案）
            wb = load_workbook(src, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — 损坏文件 / 加密表 / xls 老格式
            raise ToolError(f"打开 {src.name} 失败：{type(exc).__name__}: {exc}") from exc
        try:
            names = wb.sheetnames
            target = sheet_name if sheet_name and sheet_name.strip() else None
            if target and target not in names:
                raise ToolError(f"工作表「{target}」不存在；可用：{'、'.join(names)}")
            ws = wb[target] if target else wb[names[0]]

            rows: list[list] = []
            truncated = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= limit + 1:
                    truncated = True
                    break
                rows.append([_cell(v) for v in row])
            headers = rows[0] if rows else []
            body = rows[1:] if len(rows) > 1 else []
            return {
                "source": str(src),
                "sheet": ws.title,
                "sheets": names,
                "headers": headers,
                "rows": body,
                "row_count": len(body),
                "truncated": truncated,
            }
        finally:
            wb.close()

    return run_tool(work)


def _cell(v):
    """把单元格值转成 JSON 安全的形态。

    日期时间是最常见的「json.dumps 直接炸掉」的原因（openpyxl 给的是
    datetime 对象），所以这里统一转 ISO 串而不是让整个调用失败。
    """
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


@dsh_bridge.tool(
    name="create_styled_excel",
    description=(
        "生成一个带样式的 Excel 文件（表头加粗 + 底色 + 冻结首行 + 自动列宽 + 可选图表）。"
        "适合「把这些数据整理成一份能直接发出去的表格」。"
        "data 是一个二维数组（第一行当表头），或 {表名: 二维数组} 形式建多个工作表。"
        "返回 {path, sheets, cells_written}。"
    ),
    parameters={
        "type": "object",
        "properties": {
            "data": {
                "type": "array",
                "description": "二维数组（第一行是表头），或 {工作表名: 二维数组}",
            },
            "output_path": {"type": "string", "description": "要写出的 .xlsx 绝对路径"},
            "title": {"type": "string", "description": "可选：表头行底色标题文字（写进 A1 上方的合并标题行）"},
            "chart": {
                "type": "string",
                "description": '可选图表："bar"（首列为类目、每列一个系列）|"line"|"pie"；留空则不画',
            },
        },
        "required": ["data", "output_path"],
    },
)
def create_styled_excel(
    data,
    output_path: str,
    title: str | None = None,
    chart: str | None = None,
) -> dict:
    def work():
        try:
            import xlsxwriter
        except ImportError as exc:  # noqa: BLE001
            raise ToolError(f"XlsxWriter 不可用：{exc}。请重新执行「基本安装」。") from exc

        sheets = _normalize(data)
        dest = resolve_input_path(output_path, must_exist=False)
        if dest.suffix.lower() not in (".xlsx", ".xlsm"):
            dest = dest.with_suffix(".xlsx")
        dest.parent.mkdir(parents=True, exist_ok=True)

        wb = xlsxwriter.Workbook(str(dest))
        written = 0
        try:
            head = wb.add_format(
                {
                    "bold": True, "bg_color": "#DDEBF7", "border": 1,
                    "align": "center", "valign": "vcenter", "text_wrap": True,
                }
            )
            cell = wb.add_format({"border": 1, "valign": "top"})
            title_fmt = wb.add_format({"bold": True, "font_size": 14})
            # 数字多的列按数字格式写，长度不会顶成科学计数法
            num = wb.add_format({"border": 1, "num_format": "#,##0.00", "align": "right"})

            for name, rows in sheets.items():
                sheet_name = name[:31] or "Sheet1"
                ws = wb.add_worksheet(sheet_name)
                r0 = 0
                if title:
                    ws.merge_range(0, 0, 0, max(len(rows[0]) - 1, 0), title, title_fmt)
                    r0 = 1
                for ri, row in enumerate(rows):
                    for ci, v in enumerate(row):
                        if ri == 0:
                            ws.write(r0 + ri, ci, "" if v is None else str(v), head)
                        elif isinstance(v, (int, float)) and not isinstance(v, bool):
                            ws.write_number(r0 + ri, ci, float(v), num)
                        else:
                            ws.write(r0 + ri, ci, "" if v is None else str(v), cell)
                        written += 1
                ws.freeze_panes(r0 + 1, 0)
                for ci in range(len(rows[0])):
                    width = max((len(str(r[ci])) for r in rows if ci < len(r)), default=8)
                    ws.set_column(ci, ci, min(max(width + 2, 8), 60))
                _maybe_chart(ws, sheet_name, rows, chart, r0)

            wb.close()
        except Exception:
            # close 之前的异常会留下一个半截 xlsx；宁可删掉也不要留坏文件给用户
            wb.close()
            dest.unlink(missing_ok=True)
            raise
        return {"path": str(dest), "sheets": list(sheets.keys()), "cells_written": written}

    return run_tool(work)


def _maybe_chart(ws, sheet_name, rows, chart, r0):
    """按需画图；名字不认识或数据不够就静默跳过（画图是锦上添花，不该让整次生成失败）。"""
    kind = (chart or "").strip().lower()
    if not kind or kind not in ("bar", "line", "pie") or len(rows) < 2:
        return
    try:
        chart_obj = ws.add_chart({"type": kind})
        n_rows = len(rows)
        n_cols = len(rows[0])
        for ci in range(1, n_cols):
            # 只取「这一列真的有值」的行数，否则末尾空行会被画成 0
            last = r0 + n_rows - 1
            while last > r0 + 1 and rows[last - r0][ci] in (None, ""):
                last -= 1
            if last <= r0 + 1:
                continue
            chart_obj.add_series(
                {
                    "name": str(rows[0][ci]),
                    # 首列当类目（横轴）；数值列各自一个系列
                    "categories": [sheet_name, r0 + 1, 0, last, 0],
                    "values": [sheet_name, r0 + 1, ci, last, ci],
                }
            )
        chart_obj.set_size({"width": 520, "height": 300})
        ws.insert_chart(r0 + n_rows + 2, 0, chart_obj)
    except Exception:  # noqa: BLE001 — 图表失败不牵连数据
        return


def _normalize(data) -> dict:
    """把 data 归一成 {表名: 二维数组}。"""
    if isinstance(data, dict):
        items = list(data.items())
    elif isinstance(data, (list, tuple)):
        if not data or not isinstance(data[0], (list, tuple)):
            raise ToolError("data 必须是一个二维数组（第一行是表头），或 {表名: 二维数组}。")
        items = [("Sheet1", list(data))]
    else:
        raise ToolError("data 必须是二维数组或 {表名: 二维数组}。")
    out: dict[str, list[list]] = {}
    for name, rows in items:
        clean = [list(r) for r in rows if isinstance(r, (list, tuple))]
        if not clean:
            continue
        width = max(len(r) for r in clean)
        clean = [r + [None] * (width - len(r)) for r in clean]
        out[str(name)[:31] or "Sheet1"] = clean
    if not out:
        raise ToolError("data 里没有可写的表格。")
    return out


__all__ = ["read_excel_data", "create_styled_excel", "DEFAULT_MAX_CHARS"]